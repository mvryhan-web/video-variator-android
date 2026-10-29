from pathlib import Path
import json, math, zipfile, hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
import fitz

ROOT=Path(__file__).resolve().parent
for d in ['Buyer','Previews','Ads','Sources']: (ROOT/d).mkdir(exist_ok=True)
STYLES=[('Minimal','#F7F5F0','#242721','#687061'),('Woodland','#FAF7EF','#293E30','#65795C'),('Glamour','#132A28','#F5EBD7','#C4A76B'),('Vintage','#F2E5CF','#713D39','#9A6C50'),('Urban','#F4F1EB','#182D4C','#DA563B')]

class Page:
 def __init__(self,prs,pdf,w,h,style):
  self.slide=prs.slides.add_slide(prs.slide_layouts[6]); self.pdf=pdf; self.w=w;self.h=h
  self.name,self.bg,self.ink,self.accent=style
  self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=RGBColor.from_string(self.bg[1:])
  pdf.setFillColor(HexColor(self.bg));pdf.rect(0,0,w,h,fill=1,stroke=0)
 def line(self,x,y,x2,y2,color=None,width=.6):
  c=color or self.accent
  from pptx.enum.shapes import MSO_CONNECTOR
  s=self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x/72),Inches(y/72),Inches(x2/72),Inches(y2/72));s.line.color.rgb=RGBColor.from_string(c[1:]);s.line.width=Pt(width)
  from pptx.oxml.xmlchemy import OxmlElement
  s._element.spPr.append(OxmlElement('a:effectLst'))
  for effect in s._element.xpath('.//a:effectRef'):effect.set('idx','0')
  self.pdf.setStrokeColor(HexColor(c));self.pdf.setLineWidth(width);self.pdf.line(x,self.h-y,x2,self.h-y2)
 def text(self,text,y,size=12,font='Times-Roman',color=None,x=32,width=None,align='center'):
  width=width or self.w-2*x;c=color or self.ink
  box=self.slide.shapes.add_textbox(Inches(x/72),Inches((y-size*.2)/72),Inches(width/72),Inches(size*1.7/72));tf=box.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  p=tf.paragraphs[0];p.text=text;p.font.name='Arial' if 'Helvetica' in font else 'Times New Roman';p.font.size=Pt(size);p.font.bold='Bold' in font;p.font.italic='Italic' in font;p.font.color.rgb=RGBColor.from_string(c[1:]);p.alignment=PP_ALIGN.CENTER if align=='center' else PP_ALIGN.LEFT
  self.pdf.setFillColor(HexColor(c));self.pdf.setFont(font,size)
  if align=='center':self.pdf.drawCentredString(x+width/2,self.h-y-size*.8,text)
  else:self.pdf.drawString(x,self.h-y-size*.8,text)
 def art(self):
  w,h=self.w,self.h
  if self.name=='Woodland':
   art=ROOT/'Sources/woodland.png'
   if art.exists():
    self.slide.shapes.add_picture(str(art),0,0,width=Inches(w/72),height=Inches(h/72))
    self.pdf.drawImage(str(art),0,0,width=w,height=h)
  elif self.name=='Glamour':
   for m in [19,24]:
    self.line(m,m,w-m,m);self.line(m,h-m,w-m,h-m);self.line(m,m,m,h-m);self.line(w-m,m,w-m,h-m)
   for x in [w*.5-5,w*.5+5]:self.line(x,44,w*.5,37,width=.8);self.line(x,44,w*.5,51,width=.8)
  elif self.name=='Vintage':
   for y in [28,h-28]:
    self.line(26,y,w-26,y,width=.4);self.line(40,y+4 if y<100 else y-4,w-40,y+4 if y<100 else y-4,width=.4)
   self.text('~  *  ~',49,17,'Times-Italic',self.accent)
  elif self.name=='Urban':
   self.line(25,28,w-25,28,self.ink,3);self.line(25,h-28,w-25,h-28,self.ink,3)
   self.line(25,39,72,39,self.accent,5)
  else:
   self.line(w/2,35,w/2,66,width=.5);self.line(w/2,h-43,w/2,h-28,width=.5)

def content(pg,kind):
 h,w=pg.h,pg.w;scale=h/504;Y=lambda y:y*scale
 urban=pg.name=='Urban';head='Helvetica-Bold' if urban else 'Times-Italic';small='Helvetica'
 pg.art()
 if urban:
  if kind==0:
   for text,y,size,font,c in [('WE ARE GETTING MARRIED',85,9,'Helvetica-Bold',pg.accent),('ALEXANDRA',139,34,head,pg.ink),('& BENJAMIN',185,32,head,pg.ink),('19.06.27',260,39,head,pg.ink),('Join us for a city celebration.',328,12,'Helvetica',pg.ink),('THE GLASSHOUSE',374,12,head,pg.ink),('125 Garden Lane, Your City',401,10,'Helvetica',pg.ink),('Saturday at 4:00 PM. Dinner & dancing to follow.',434,8.5,'Helvetica',pg.ink)]:pg.text(text,Y(y),size,font,c,x=27,width=w-54,align='left')
  elif kind==1:
   pg.text('THE PLAN',Y(94),36,head,x=27,align='left')
   for y,time,event in [(178,'16:00','The ceremony'),(232,'17:00','Cocktails & conversation'),(286,'18:30','Dinner together'),(340,'20:30','Cake. Music. Dance floor.')]:
    pg.text(time,Y(y),18,head,pg.accent,x=27,align='left');pg.text(event,Y(y+24),12,'Helvetica',x=27,align='left')
   pg.text('DRESS: COCKTAIL ATTIRE',Y(412),9,head,x=27,align='left');pg.text('Travel: your-wedding-website.com',Y(439),10,'Helvetica',x=27,align='left')
  else:
   pg.text('ARE YOU IN?',Y(110),34,head,x=27,align='left');pg.text('Please reply by 19 May 2027.',Y(177),12,'Helvetica',x=27,align='left')
   for y,text in [(237,'YOUR NAME(S)'),(312,'ATTENDING / UNABLE TO ATTEND'),(377,'DIETARY REQUIREMENTS')]:
    pg.text(text,Y(y),9,head,x=27,align='left');pg.line(27,Y(y+40),w-27,Y(y+40),width=.6)
   pg.text('your-wedding-website.com',Y(443),10,'Helvetica',x=27,align='left')
  return
 if kind==0:
  opening={'Minimal':'WITH FULL HEARTS','Woodland':'A LOVE ROOTED IN SOMETHING BEAUTIFUL','Glamour':'AN EVENING TO REMEMBER','Vintage':'A NEW CHAPTER BEGINS'}[pg.name]
  pg.text(opening,Y(123 if pg.name=='Woodland' else 102),7.5,small)
  pg.text('Alexandra',Y(151),36,head)
  pg.text('&',Y(195),18,'Times-Italic',pg.accent)
  pg.text('Benjamin',Y(225),36,head)
  phrase={'Minimal':'invite you to celebrate their wedding','Woodland':'invite you to a wedding in the garden','Glamour':'request the pleasure of your company','Vintage':'invite you to share their happily ever after'}[pg.name]
  pg.text(phrase,Y(289),10.5,'Times-Roman')
  pg.text('SATURDAY  |  19 JUNE 2027',Y(331),10,'Helvetica-Bold')
  pg.text('CEREMONY AT FOUR IN THE AFTERNOON',Y(354),7.5,small)
  pg.text('The Glasshouse',Y(384),17,'Times-Italic')
  pg.text('125 Garden Lane, Your City',Y(410),10)
  pg.text('Dinner, dancing & a beautiful beginning',Y(423 if pg.name=='Woodland' else 438),9 if pg.name=='Woodland' else 10,'Times-Italic')
 elif kind==1:
  pg.text('THE CELEBRATION',Y(95),8,small)
  pg.text('The details',Y(132),34,head)
  for y,label,value in [(211,'THE DAY','4:00 PM Ceremony  /  5:00 PM Drinks'),(269,'THE EVENING','6:30 PM Dinner  /  8:30 PM Cake & dancing'),(327,'DRESS CODE','Cocktail attire; garden-friendly shoes'),(385,'TRAVEL & STAY','Visit our wedding website for local stays')]:
   pg.text(label,Y(y),8,'Helvetica-Bold',pg.accent)
   pg.text(value,Y(y+22),10)
  if pg.name=='Woodland':pg.text('your-wedding-website.com',Y(442),9,'Times-Italic',x=32,width=205,align='left')
  else:pg.text('your-wedding-website.com',Y(442),11,'Times-Italic')
 else:
  pg.text('ALEXANDRA & BENJAMIN',Y(102),8,small)
  pg.text('Kindly reply',Y(151),34,head)
  pg.text('Please respond by 19 May 2027',Y(215),12)
  pg.text('Name(s)',Y(265),10,'Times-Italic')
  pg.line(58,Y(296),w-58,Y(296))
  pg.text('Joyfully accepts  /  Regretfully declines',Y(323),10)
  pg.text('Dietary requirements',Y(365),10,'Times-Italic')
  pg.line(58,Y(396),w-58,Y(396))
  if pg.name=='Woodland':
   pg.text('Reply through our wedding website',Y(419),9,x=32,width=205,align='left')
   pg.text('your-wedding-website.com',Y(442),9,'Times-Italic',x=32,width=205,align='left')
  else:
   pg.text('Reply through our wedding website',Y(430),10)
   pg.text('your-wedding-website.com',Y(451),11,'Times-Italic')

manifest=[]
for style in STYLES:
 for mode in ['Print','Mobile']:
  w,h=(360,504) if mode=='Print' else (360,640)
  prs=Presentation();prs.slide_width=Inches(w/72);prs.slide_height=Inches(h/72)
  stem=f'{style[0]}_{mode}';out=ROOT/'Buyer'/f'{stem}_Sample.pdf';tmp_pdf=Path('/tmp')/out.name;pdf=canvas.Canvas(str(tmp_pdf),pagesize=(w,h));pdf.setTitle(stem+' - editable sample');pdf.setAuthor('Pinogri')
  for i in range(3):
   content(Page(prs,pdf,w,h,style),i);pdf.showPage()
  pdf.save();tmp_pdf.replace(out);tmp_pptx=Path('/tmp')/f'{stem}.pptx';prs.save(tmp_pptx);tmp_pptx.replace(ROOT/'Buyer'/f'{stem}.pptx')
  doc=fitz.open(out)
  for i,page in enumerate(doc):
   dest=ROOT/'Previews'/f'{stem}_{i+1:02}.png';tmp=Path('/tmp')/dest.name;page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(tmp);tmp.replace(dest)
   manifest.append({'style':style[0],'version':mode,'kind':['Invitation','Details','RSVP'][i],'size_inches':[w/72,h/72],'preview':str(dest.relative_to(ROOT))})
(ROOT/'manifest.json').write_text(json.dumps(manifest,indent=2))
print('Built',len(manifest),'pages; 15 distinct layouts, 15 mobile adaptations. PDFs and editable PPTX.')
